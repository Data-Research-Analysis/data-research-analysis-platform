#!/usr/bin/env python3
"""Builds a concise client-facing digital marketing strategy deck (PPTX) from a structured content JSON.

Usage:
    python build_deck.py <content.json> [output.pptx]

Defaults:
    output.pptx  -> <dir of content.json>/00-digital-marketing-strategy-deck.pptx

Content JSON schema (see SKILL.md for how to fill it from the plan pack):

{
  "meta": {
    "client": "str",            # company name
    "campaign": "str",          # campaign / test name
    "date": "str",              # presentation date
    "prepared_by": "str",       # agency / team name
    "platforms": "str",         # e.g. Google Search | Google Display Network | Meta Ads
    "budget": "str",            # e.g. $250,000 USD
    "markets": "str",           # e.g. United States (national)
    "window": "str"             # e.g. 2026-10-19 to 2026-11-17 (30 days)
  },
  "slides": [
    {"type": "title", "kicker": "str", "title": "str", "subtitle": "str"},
    {"type": "agenda", "title": "str", "items": [
        {"num": "01", "heading": "str", "desc": "str"}, ... ]},
    {"type": "section", "number": "01", "title": "str", "subtitle": "str"},
    {"type": "execsummary", "title": "str", "kicker": "str", "paragraph": "str",
     "stats": [{"value": "str", "label": "str"}, ...]},        # 2-4 stats
    {"type": "bullets", "title": "str", "kicker": "str", "lead": "str",
     "bullets": ["str", ...]},                                  # <= 6 bullets
    {"type": "table", "title": "str", "kicker": "str",
     "header": ["str", ...], "rows": [["str", ...], ...],
     "footnote": "str", "align": ["l", "r", "c", ...],
     "col_widths": [2, 2, 1, 4, 1]},                # optional relative weights
    {"type": "statgrid", "title": "str", "kicker": "str",
     "stats": [{"value": "str", "label": "str", "note": "str"}, ...]},  # <= 6
    {"type": "twocol", "title": "str", "kicker": "str",
     "left": {"heading": "str", "bullets": ["str", ...]},
     "right": {"heading": "str", "bullets": ["str", ...]}},
    {"type": "companypaid", "title": "str", "kicker": "str", "lead": "str",
     "header": ["Platform", "Verdict", "What the company is running"],
     "rows": [["str", "str", "str"], ...],          # omit/empty -> states "no verifiable info"
     "footnote": "str", "note": "str"},             # note used when no rows
    {"type": "nextsteps", "title": "str", "kicker": "str",
     "items": ["str", ...]},                                     # <= 6
    {"type": "closing", "title": "str", "subtitle": "str"}
  ]
}

Requires: python-pptx (pip install python-pptx). 16:9 slides (13.333 x 7.5 in).
"""
import argparse
import json
import math
import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

# ----------------------------------------------------------------------------
# Theme
# ----------------------------------------------------------------------------
NAVY = RGBColor(0x0F, 0x2D, 0x52)
INK = RGBColor(0x1F, 0x29, 0x33)
SUB = RGBColor(0x5B, 0x6B, 0x7C)
TEAL = RGBColor(0x0E, 0x7C, 0x86)
AMBER = RGBColor(0xE8, 0xA3, 0x3D)
CARD = RGBColor(0xF1, 0xF5, 0xF9)
CARD2 = RGBColor(0xF8, 0xFA, 0xFC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xE2, 0xE8, 0xF0)
ONNAVY = RGBColor(0xCF, 0xDB, 0xEA)
FONT = "Calibri"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = 0.6
CONTENT_W = 13.333 - 2 * MARGIN  # 12.133 in
BODY_TOP = 1.55

FOOTER_TEXT = (
    "Digital Strategy Built By Data Research Analysis "
    "(SMC-Private) Limited (www.dataresearchanalysis.com)"
)
DARK_SLIDE_TYPES = {"title", "section", "closing"}


# ----------------------------------------------------------------------------
# Low-level helpers
# ----------------------------------------------------------------------------
def warn_overflow(label, text, size_pt, width_in, height_in):
    cpl = max(1, int((width_in * 72) / (size_pt * 0.55)))
    lines = 0
    for seg in str(text).split("\n"):
        lines += max(1, math.ceil(len(seg) / cpl))
    line_h = size_pt * 1.25 / 72
    need = lines * line_h
    if need > height_in:
        print(
            f"[WARN] overflow risk on '{label}': ~{lines} lines need "
            f"{need:.2f}in, box is {height_in:.2f}in"
        )


def set_para(para, text, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT,
             space_after=6, space_before=0, italic=False, font=FONT, line_spacing=None):
    para.text = text
    for run in para.runs:
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
        run.font.name = font
    para.alignment = align
    para.space_after = Pt(space_after)
    para.space_before = Pt(space_before)
    if line_spacing:
        para.line_spacing = line_spacing
    return para


def add_text(slide, left, top, width, height, text="", size=14, bold=False,
             color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             wrap=True, space_after=6, italic=False, font=FONT):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    set_para(p, text, size=size, bold=bold, color=color, align=align,
             space_after=space_after, italic=italic, font=font)
    return tb, tf


def add_rect(slide, left, top, width, height, fill=WHITE, line=None,
             shape=MSO_SHAPE.RECTANGLE, radius=None):
    shp = slide.shapes.add_shape(shape, left, top, width, height)
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    shp.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    return shp


def add_oval_text(slide, left, top, width, height, text, size, bold, color, fill):
    shp = add_rect(slide, left, top, width, height, fill=fill, shape=MSO_SHAPE.OVAL)
    tf = shp.text_frame
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    set_para(p, text, size=size, bold=bold, color=color,
             align=PP_ALIGN.CENTER, space_after=0)
    return shp


def content_header(slide, kicker, title, page, footer=True):
    add_rect(slide, 0, 0, SLIDE_W, Inches(0.09), fill=NAVY)
    add_text(slide, Inches(MARGIN), Inches(0.30), Inches(CONTENT_W), Inches(0.28),
             kicker.upper(), 11, True, TEAL, space_after=0)
    add_text(slide, Inches(MARGIN), Inches(0.58), Inches(CONTENT_W), Inches(0.66),
             title, 26, True, NAVY, space_after=0)
    add_rect(slide, Inches(MARGIN + 0.02), Inches(1.28), Inches(0.9), Inches(0.05), fill=AMBER)
    if footer:
        add_rect(slide, Inches(MARGIN), SLIDE_H - Inches(0.55),
                 Inches(CONTENT_W), Inches(0.014), fill=LIGHT)
        add_text(slide, SLIDE_W - Inches(0.85), SLIDE_H - Inches(0.46),
                 Inches(0.5), Inches(0.3), str(page), 10, False, SUB,
                 align=PP_ALIGN.RIGHT, space_after=0)


def add_slide_footer(slide, slide_type):
    """Stamp the agency attribution on every slide."""
    color = ONNAVY if slide_type in DARK_SLIDE_TYPES else SUB
    add_text(
        slide,
        Inches(MARGIN),
        SLIDE_H - Inches(0.46),
        Inches(CONTENT_W - 0.9),
        Inches(0.3),
        FOOTER_TEXT,
        9,
        False,
        color,
        align=PP_ALIGN.LEFT,
        space_after=0,
    )


def stat_cards(slide, stats, top, width_per, height, value_size=30, max_cols=6):
    n = min(len(stats), max_cols)
    gap = 0.3
    w = width_per
    for i in range(n):
        left = Inches(MARGIN + i * (w + gap))
        card = add_rect(slide, left, Inches(top), Inches(w), Inches(height),
                        fill=CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
        tf = card.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Inches(0.15)
        tf.margin_right = Inches(0.15)
        tf.margin_top = Inches(0.08)
        tf.margin_bottom = Inches(0.08)
        p = tf.paragraphs[0]
        set_para(p, stats[i]["value"], size=value_size, bold=True, color=NAVY,
                 align=PP_ALIGN.CENTER, space_after=2)
        p2 = tf.add_paragraph()
        set_para(p2, stats[i]["label"], size=12, bold=False, color=SUB,
                 align=PP_ALIGN.CENTER, space_after=0)
        warn_overflow("stat-" + stats[i]["label"], stats[i]["value"], value_size, w, 0.8)
    if len(stats) > max_cols:
        print(f"[WARN] statgrid capped at {max_cols} cards; {len(stats)} provided")


def bullet_block(slide, bullets, left, top, width, height, size=15, gap_pt=10):
    y = top
    for i, b in enumerate(bullets):
        add_rect(slide, Inches(left), Inches(y + 0.07), Inches(0.11), Inches(0.11), fill=TEAL)
        _, tf = add_text(slide, Inches(left + 0.25), Inches(y), Inches(width - 0.25), Inches(0.9),
                         b, size=size, color=INK, space_after=0, wrap=True)
        warn_overflow(f"bullet-{i}", b, size, width - 0.25, 0.75)
        tf.paragraphs[0].line_spacing = 1.05
        y += size * 1.05 / 72 * 2.6 + gap_pt / 72
    return y


# ----------------------------------------------------------------------------
# Slide renderers
# ----------------------------------------------------------------------------
def render_title(prs, s, page):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, fill=NAVY)
    add_rect(slide, 0, Inches(0.0), Inches(0.18), SLIDE_H, fill=AMBER)
    add_text(slide, Inches(1.0), Inches(1.7), Inches(11.0), Inches(0.4),
             s.get("kicker", "DIGITAL MARKETING STRATEGY").upper(),
             14, True, AMBER, space_after=0)
    add_text(slide, Inches(1.0), Inches(2.15), Inches(11.3), Inches(2.0),
             s.get("title", ""), 40, True, WHITE, space_after=0)
    add_rect(slide, Inches(1.02), Inches(3.95), Inches(2.0), Inches(0.07), fill=AMBER)
    add_text(slide, Inches(1.0), Inches(4.2), Inches(11.3), Inches(1.0),
             s.get("subtitle", ""), 18, False, ONNAVY, space_after=0)

    chips = [
        ("Budget", prs.meta.get("budget", "")),
        ("Markets", prs.meta.get("markets", "")),
        ("Platforms", prs.meta.get("platforms", "")),
        ("Window", prs.meta.get("window", "")),
    ]
    w = (11.33 - 3 * 0.35) / 4
    for i, (label, value) in enumerate(chips):
        left = Inches(1.0 + i * (w + 0.35))
        card = add_rect(slide, left, Inches(5.7), Inches(w), Inches(0.85),
                        fill=RGBColor(0x17, 0x3A, 0x66),
                        shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
        tf = card.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Inches(0.12)
        tf.margin_right = Inches(0.12)
        p = tf.paragraphs[0]
        set_para(p, label.upper(), size=9, bold=True, color=AMBER,
                 align=PP_ALIGN.CENTER, space_after=1)
        p2 = tf.add_paragraph()
        set_para(p2, value, size=12, bold=False, color=WHITE,
                 align=PP_ALIGN.CENTER, space_after=0)
        warn_overflow("chip-" + label, value, 12, w - 0.24, 0.5)
    return slide


def render_agenda(prs, s, page):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    content_header(slide, s.get("kicker", "AGENDA"), s.get("title", "Agenda"), page)
    items = s.get("items", [])
    col_w = (CONTENT_W - 0.5) / 2
    for i, item in enumerate(items):
        col = i // 4
        row = i % 4
        if isinstance(item, dict):
            num = item.get("num", "")
            heading = item.get("heading", "")
            desc = item.get("desc", "")
        else:
            num = str(i + 1).zfill(2)
            heading = item
            desc = ""
        left = Inches(MARGIN + col * (col_w + 0.5))
        top = Inches(BODY_TOP + row * 1.3)
        add_oval_text(slide, left, top, Inches(0.55), Inches(0.55),
                      num, 15, True, WHITE, NAVY)
        add_text(slide, left + Inches(0.75), top, Inches(col_w - 0.8), Inches(0.4),
                 heading, 16, True, INK, space_after=0)
        if desc:
            add_text(slide, left + Inches(0.75), top + Inches(0.4), Inches(col_w - 0.8), Inches(0.7),
                     desc, 12, False, SUB, space_after=0, wrap=True)
            warn_overflow("agenda-" + heading, desc, 12, col_w - 0.8, 0.7)
    return slide


def render_section(prs, s, page):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, fill=NAVY)
    add_rect(slide, 0, 0, Inches(0.18), SLIDE_H, fill=AMBER)
    add_text(slide, Inches(1.0), Inches(2.0), Inches(3.5), Inches(1.6),
             s.get("number", ""), 72, True, AMBER, space_after=0)
    add_text(slide, Inches(1.05), Inches(3.7), Inches(11.3), Inches(1.0),
             s.get("title", ""), 34, True, WHITE, space_after=0)
    add_rect(slide, Inches(1.07), Inches(4.6), Inches(2.0), Inches(0.06), fill=AMBER)
    add_text(slide, Inches(1.05), Inches(4.85), Inches(11.0), Inches(1.2),
             s.get("subtitle", ""), 16, False, ONNAVY, space_after=0, wrap=True)
    return slide


def render_execsummary(prs, s, page):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    content_header(slide, s.get("kicker", "EXECUTIVE SUMMARY"), s.get("title", "Executive Summary"), page)
    stats = s.get("stats", [])
    stat_cards(slide, stats, BODY_TOP, 2.81, 1.3, value_size=28, max_cols=4)
    para = s.get("paragraph", "")
    _, tf = add_text(slide, Inches(MARGIN), Inches(3.25), Inches(CONTENT_W), Inches(3.4),
                     para, size=16, color=INK, space_after=0, wrap=True)
    tf.paragraphs[0].line_spacing = 1.25
    warn_overflow("exec-summary", para, 16, CONTENT_W, 3.4)
    return slide


def render_bullets(prs, s, page):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    content_header(slide, s.get("kicker", "FINDINGS"), s.get("title", ""), page)
    lead = s.get("lead", "")
    top = BODY_TOP
    if lead:
        _, tf = add_text(slide, Inches(MARGIN), Inches(top), Inches(CONTENT_W), Inches(0.9),
                         lead, size=15, color=SUB, italic=True, space_after=0, wrap=True)
        tf.paragraphs[0].line_spacing = 1.1
        warn_overflow("lead", lead, 15, CONTENT_W, 0.9)
        top += 1.05
    bullets = s.get("bullets", [])
    bullet_block(slide, bullets, MARGIN, top, CONTENT_W, 4.5, size=15, gap_pt=8)
    return slide


def draw_table(slide, header, rows, align=None, col_widths=None, footnote="", top=BODY_TOP):
    header = header or []
    rows = rows or []
    align = align or ["l"] * len(header)
    ncols = len(header) or (len(rows[0]) if rows else 0)
    nrows = len(rows) + 1
    height = 0.38 + 0.5 * len(rows)
    if height > 4.6:
        height = 4.6
    gfx = slide.shapes.add_table(nrows, ncols, Inches(MARGIN), Inches(top),
                                 Inches(CONTENT_W), Inches(height))
    table = gfx.table
    table.first_row = False
    table.last_row = False
    table.first_col = False
    table.last_col = False
    table.horz_banding = False
    table.vert_banding = False
    if col_widths and len(col_widths) == ncols:
        total_w = sum(float(w) for w in col_widths)
        col_w = [CONTENT_W * float(w) / total_w for w in col_widths]
    else:
        col_w = [CONTENT_W / ncols] * ncols
    for c in range(ncols):
        table.columns[c].width = Inches(col_w[c])
    table.rows[0].height = Inches(0.42)
    for r in range(1, nrows):
        table.rows[r].height = Inches(0.46)

    h_size = 11 if ncols > 5 else 12
    b_size = 10 if ncols > 5 else 11
    for c, cell_text in enumerate(header):
        cell = table.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Inches(0.06)
        cell.margin_right = Inches(0.06)
        set_para(cell.text_frame.paragraphs[0], cell_text, size=h_size, bold=True,
                 color=WHITE, align=PP_ALIGN.CENTER, space_after=0)
    for r, row in enumerate(rows):
        fill = CARD if r % 2 == 0 else WHITE
        for c, cell_text in enumerate(row):
            cell = table.cell(r + 1, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = Inches(0.06)
            cell.margin_right = Inches(0.06)
            align_ = PP_ALIGN.LEFT if align[c] == "l" else (PP_ALIGN.CENTER if align[c] == "c" else PP_ALIGN.RIGHT)
            set_para(cell.text_frame.paragraphs[0], cell_text, size=b_size, bold=False,
                     color=INK, align=align_, space_after=0)
            warn_overflow(f"tbl-{r}-{c}", cell_text, b_size, col_w[c] - 0.12, 0.5)
    if footnote:
        add_text(slide, Inches(MARGIN), Inches(top + height + 0.15), Inches(CONTENT_W), Inches(0.4),
                 footnote, size=10, color=SUB, italic=True, space_after=0, wrap=True)
    return top + height


def render_table(prs, s, page):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    content_header(slide, s.get("kicker", "DATA"), s.get("title", ""), page)
    draw_table(slide, s.get("header", []), s.get("rows", []),
               align=s.get("align"), col_widths=s.get("col_widths"),
               footnote=s.get("footnote", ""))
    return slide


def render_companypaid(prs, s, page):
    """One slide on the company's own paid media activity (Step 3A audit).

    Content keys:
      kicker, title, lead (optional), rows (list of [platform, verdict, detail]),
      header (optional override), align, col_widths, footnote,
      note (used when there are no rows / nothing verifiable).
    If no verifiable activity is supplied, the slide states that explicitly
    instead of omitting the section.
    """
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    content_header(slide, s.get("kicker", "Company Paid Media Audit"),
                   s.get("title", "What the Company Is Running Today"), page)
    lead = s.get("lead", "")
    top = BODY_TOP
    if lead:
        _, tf = add_text(slide, Inches(MARGIN), Inches(top), Inches(CONTENT_W), Inches(0.9),
                         lead, size=15, color=SUB, italic=True, space_after=0, wrap=True)
        tf.paragraphs[0].line_spacing = 1.1
        warn_overflow("cp-lead", lead, 15, CONTENT_W, 0.9)
        top += 1.0
    rows = s.get("rows", [])
    if rows:
        header = s.get("header", ["Platform", "Verdict", "What the company is running"])
        draw_table(slide, header, rows,
                   align=s.get("align", ["l"] * len(header)),
                   col_widths=s.get("col_widths"),
                   footnote=s.get("footnote", ""), top=top)
    else:
        note = s.get("note", "No verifiable information was found for the company's own paid media "
                             "activity on the audited platforms. Ad-library and transparency data "
                             "could not be retrieved; this is reported as NOT VERIFIED rather than "
                             "assumed to be inactive.")
        box = add_rect(slide, Inches(MARGIN), Inches(top + 0.2), Inches(CONTENT_W), Inches(1.8),
                       fill=CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
        tf = box.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Inches(0.2)
        tf.margin_right = Inches(0.2)
        set_para(tf.paragraphs[0], note, size=15, color=INK, align=PP_ALIGN.LEFT, space_after=0)
        warn_overflow("cp-note", note, 15, CONTENT_W - 0.4, 1.6)
    return slide


def render_statgrid(prs, s, page):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    content_header(slide, s.get("kicker", "KEY NUMBERS"), s.get("title", ""), page)
    stats = s.get("stats", [])
    n = min(len(stats), 6)
    cols = 3 if n > 3 else n
    rows_n = math.ceil(n / cols)
    gap = 0.3
    w = (CONTENT_W - (cols - 1) * gap) / cols
    h = 1.55 if rows_n > 1 else 1.8
    for i, st in enumerate(stats[:6]):
        col = i % cols
        row = i // cols
        left = Inches(MARGIN + col * (w + gap))
        top = Inches(BODY_TOP + row * (h + 0.35))
        card = add_rect(slide, left, top, Inches(w), Inches(h),
                        fill=CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
        tf = card.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = Inches(0.15)
        tf.margin_right = Inches(0.15)
        tf.margin_top = Inches(0.12)
        p = tf.paragraphs[0]
        set_para(p, st.get("value", ""), size=26, bold=True, color=NAVY,
                 align=PP_ALIGN.LEFT, space_after=3)
        p2 = tf.add_paragraph()
        set_para(p2, st.get("label", ""), size=13, bold=True, color=INK,
                 align=PP_ALIGN.LEFT, space_after=2)
        note = st.get("note", "")
        if note:
            p3 = tf.add_paragraph()
            set_para(p3, note, size=10, color=SUB, italic=True,
                     align=PP_ALIGN.LEFT, space_after=0)
            warn_overflow(f"grid-{st.get('label')}", note, 10, w - 0.3, 0.5)
    return slide


def render_twocol(prs, s, page):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    content_header(slide, s.get("kicker", "DETAIL"), s.get("title", ""), page)
    col_w = (CONTENT_W - 0.5) / 2
    for i, side in enumerate(("left", "right")):
        block = s.get(side, {})
        left = Inches(MARGIN + i * (col_w + 0.5))
        add_text(slide, left, Inches(BODY_TOP), Inches(col_w), Inches(0.4),
                 block.get("heading", "").upper(), 13, True, TEAL, space_after=0)
        bullet_block(slide, block.get("bullets", []), MARGIN + i * (col_w + 0.5),
                     BODY_TOP + 0.5, col_w, 4.3, size=14, gap_pt=7)
    return slide


def render_nextsteps(prs, s, page):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    content_header(slide, s.get("kicker", "NEXT ACTIONS"), s.get("title", ""), page)
    items = s.get("items", [])
    top = BODY_TOP
    step_h = 0.82
    for i, item in enumerate(items):
        add_oval_text(slide, Inches(MARGIN), Inches(top + 0.08), Inches(0.5), Inches(0.5),
                      str(i + 1), 14, True, WHITE, TEAL)
        add_rect(slide, Inches(MARGIN + 0.72), Inches(top), Inches(CONTENT_W - 0.72), Inches(step_h - 0.1),
                 fill=CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
        add_text(slide, Inches(MARGIN + 0.95), Inches(top + 0.1), Inches(CONTENT_W - 1.25), Inches(step_h - 0.2),
                 item, size=14, color=INK, space_after=0, wrap=True)
        warn_overflow(f"step-{i}", item, 14, CONTENT_W - 1.25, step_h - 0.2)
        top += step_h + 0.1
    return slide


def render_closing(prs, s, page):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, fill=NAVY)
    add_rect(slide, 0, 0, Inches(0.18), SLIDE_H, fill=AMBER)
    add_text(slide, Inches(1.0), Inches(2.4), Inches(11.3), Inches(1.2),
             s.get("title", "Thank you"), 40, True, WHITE, space_after=0)
    add_text(slide, Inches(1.02), Inches(3.5), Inches(11.0), Inches(1.6),
             s.get("subtitle", ""), 18, False, ONNAVY, space_after=0, wrap=True)
    add_text(slide, Inches(1.0), Inches(6.6), Inches(11.3), Inches(0.5),
             f"Prepared by {prs.meta.get('prepared_by', '')}  ·  {prs.meta.get('date', '')}",
             12, False, AMBER, space_after=0)
    return slide


RENDERERS = {
    "title": render_title,
    "agenda": render_agenda,
    "section": render_section,
    "execsummary": render_execsummary,
    "bullets": render_bullets,
    "table": render_table,
    "companypaid": render_companypaid,
    "statgrid": render_statgrid,
    "twocol": render_twocol,
    "nextsteps": render_nextsteps,
    "closing": render_closing,
}


# ----------------------------------------------------------------------------
# Main
# ----------------------------------------------------------------------------
def check_bounds(prs):
    """Fail the build if any shape sits outside the slide canvas (catch misplaced text)."""
    w_emu = prs.slide_width
    h_emu = prs.slide_height
    tol = 5000  # EMU tolerance (~0.005 in)
    problems = []
    for i, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if sh.left is None or sh.top is None or sh.width is None or sh.height is None:
                continue
            l, t, r, b = sh.left, sh.top, sh.left + sh.width, sh.top + sh.height
            if l < -tol or t < -tol or r > w_emu + tol or b > h_emu + tol:
                label = sh.text_frame.text.strip().split("\n")[0][:40] if sh.has_text_frame and sh.text_frame.text.strip() else str(sh.shape_type)
                problems.append(f"  slide {i}: '{label}' -> left {l} top {t} right {r} bottom {b}")
    if problems:
        print("[ERROR] off-canvas shapes found (text would be invisible):")
        for p in problems:
            print(p)
        print("[ERROR] fix the content/layout and rebuild. Build aborted.")
        sys.exit(1)
    print("[OK] all shapes within slide bounds")


def main():
    ap = argparse.ArgumentParser(description="Build the client marketing strategy deck")
    ap.add_argument("content", help="Path to the deck content JSON")
    ap.add_argument("output", nargs="?", help="Output PPTX path (default: alongside content)")
    args = ap.parse_args()

    with open(args.content, "r", encoding="utf-8") as fh:
        data = json.load(fh)

    meta = data.get("meta", {})
    out = args.output or os.path.join(
        os.path.dirname(os.path.abspath(args.content)),
        "00-digital-marketing-strategy-deck.pptx",
    )

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    prs.meta = meta  # expose to renderers

    page = 0
    for s in data.get("slides", []):
        s_type = s.get("type", "bullets")
        page += 1
        renderer = RENDERERS.get(s_type)
        if not renderer:
            print(f"[WARN] unknown slide type '{s_type}', skipped")
            continue
        try:
            slide = renderer(prs, s, page)
            add_slide_footer(slide, s_type)
        except Exception as exc:
            print(f"[ERROR] slide {page} ({s_type}): {exc}")
            sys.exit(1)

    prs.save(out)
    verify = Presentation(out)
    check_bounds(verify)
    n = len(verify.slides._sldIdLst)
    print(f"[OK] wrote {out} with {n} slides")
    print("[DONE] deck built. Review overflow warnings above; if any, shorten the text and rebuild.")


if __name__ == "__main__":
    main()
